from __future__ import annotations

from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from csv2json.converter import convert_csv_text


@dataclass(frozen=True)
class IngestOptions:
    format: str = "auto"  # auto|csv|tsv|xlsx|pdf
    delimiter: str = ","
    encoding: str = "utf-8"

    # PDF-specific
    # - auto: prefer tables; if none found, fall back to text
    # - tables: tables only
    # - text: text only
    # - both: return both tables + text
    pdf_mode: str = "auto"  # auto|tables|text|both
    max_pdf_pages: int = 25
    max_pdf_chars: int = 2_000_000


def detect_format(filename: str | None) -> str:
    if not filename:
        return "csv"

    suffix = Path(filename).suffix.lower().lstrip(".")
    if suffix in {"csv"}:
        return "csv"
    if suffix in {"tsv"}:
        return "tsv"
    if suffix in {"xlsx"}:
        return "xlsx"
    if suffix in {"pdf"}:
        return "pdf"

    # Default: treat unknown as CSV.
    return "csv"


def _coerce_header(header: Iterable[Any]) -> list[str]:
    out: list[str] = []
    for i, h in enumerate(header):
        if h is None or str(h).strip() == "":
            out.append(f"col_{i+1}")
        else:
            out.append(str(h).strip())
    return out


def ingest_file(
    path: str | Path,
    *,
    filename: str | None = None,
    options: IngestOptions,
) -> Any:
    fmt = options.format
    if fmt == "auto":
        fmt = detect_format(filename or str(path))

    if fmt == "tsv":
        return ingest_csv(
            path,
            options=IngestOptions(
                format=options.format,
                delimiter="\t",
                encoding=options.encoding,
                pdf_mode=options.pdf_mode,
                max_pdf_pages=options.max_pdf_pages,
                max_pdf_chars=options.max_pdf_chars,
            ),
        )

    if fmt == "csv":
        return ingest_csv(path, options=options)

    if fmt == "xlsx":
        return ingest_xlsx(path)

    if fmt == "pdf":
        return ingest_pdf(
            path,
            max_pages=options.max_pdf_pages,
            mode=options.pdf_mode,
            max_chars=options.max_pdf_chars,
        )

    raise ValueError(f"Unsupported format: {fmt}")


def ingest_csv(path: str | Path, *, options: IngestOptions) -> list[dict[str, Any]]:
    with open(path, encoding=options.encoding, newline="") as fp:
        return convert_csv_text(fp, delimiter=options.delimiter)


def ingest_xlsx(path: str | Path) -> list[dict[str, Any]]:
    # Keep this import local so non-web usage doesn't require openpyxl.
    from openpyxl import load_workbook

    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    ws = wb.active

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration as exc:
        raise ValueError("XLSX has no rows") from exc

    header = _coerce_header(header_row)

    out: list[dict[str, Any]] = []
    for row in rows_iter:
        if row is None:
            continue
        item: dict[str, Any] = {}
        for i, key in enumerate(header):
            val = row[i] if i < len(row) else None
            item[key] = val
        out.append(item)

    return out


def ingest_pdf(
    path: str | Path,
    *,
    max_pages: int = 25,
    mode: str = "auto",
    max_chars: int = 2_000_000,
) -> dict[str, Any]:
    # Keep this import local so non-web usage doesn't require pdfplumber.
    import pdfplumber

    # We support:
    # - tables extraction (page.extract_tables)
    # - text extraction (page.extract_text)
    rows: list[dict[str, Any]] = []
    pages: list[dict[str, Any]] = []

    with pdfplumber.open(str(path)) as pdf:
        if len(pdf.pages) == 0:
            raise ValueError("PDF has no pages")

        page_count = min(len(pdf.pages), max_pages)

        def extract_tables() -> None:
            nonlocal rows
            for page_index in range(page_count):
                page = pdf.pages[page_index]
                tables = page.extract_tables() or []
                for table in tables:
                    if not table or len(table) < 2:
                        continue
                    header = _coerce_header(table[0])
                    for data_row in table[1:]:
                        item: dict[str, Any] = {}
                        for i, key in enumerate(header):
                            val = data_row[i] if i < len(data_row) else None
                            item[key] = val
                        rows.append(item)

        def extract_text() -> None:
            nonlocal pages
            used = 0
            truncated = False

            for page_index in range(page_count):
                page = pdf.pages[page_index]
                text = page.extract_text() or ""
                # Enforce a global character limit to avoid huge responses.
                remaining = max_chars - used
                if remaining <= 0:
                    truncated = True
                    break

                if len(text) > remaining:
                    text = text[:remaining]
                    truncated = True

                used += len(text)
                pages.append({"page": page_index + 1, "text": text})

            if truncated and pages:
                pages.append(
                    {
                        "page": None,
                        "text": "[truncated: PDF text exceeded max_pdf_chars]",
                    }
                )

        mode_norm = mode.lower().strip()
        if mode_norm not in {"auto", "tables", "text", "both"}:
            raise ValueError("pdf_mode must be one of: auto, tables, text, both")

        if mode_norm in {"tables", "both", "auto"}:
            extract_tables()

        if mode_norm in {"text", "both"}:
            extract_text()

        if mode_norm == "auto" and not rows:
            extract_text()

    if mode_norm == "tables" and not rows:
        raise ValueError("No tables found in PDF")

    if mode_norm == "text" and not any(p.get("text") for p in pages):
        raise ValueError("No extractable text found in PDF")

    if mode_norm == "both":
        return {"kind": "pdf", "tables": rows, "pages": pages}

    if rows and mode_norm in {"auto", "tables"}:
        return {"kind": "pdf_tables", "rows": rows}

    return {"kind": "pdf_text", "pages": pages}
